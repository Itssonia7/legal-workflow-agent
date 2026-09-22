import os
import hashlib
from rest_framework import views, status, permissions
from rest_framework.response import Response
from django.shortcuts import get_object_or_404
from django.conf import settings

from .models import CaseFile, LegalDocument
from .serializers import LegalDocumentSerializer

# Import AI Engine components
from ai_engine.ingest import process_legal_document
from ai_engine.graph import app as ai_app

class DocumentUploadAndIngestView(views.APIView):
    """
    POST /api/cases/documents/upload/
    Uploads a legal PDF document, saves it to the database,
    and runs the local ChromaDB ingestion pipeline.
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, *args, **kwargs):
        case_id = request.data.get('case_file')
        file_obj = request.FILES.get('file')

        if not case_id or not file_obj:
            return Response(
                {"error": "Both 'case_file' ID and 'file' are required."},
                status=status.HTTP_400_BAD_REQUEST
            )

        case_file = get_object_or_404(CaseFile, id=case_id, lawyer=request.user)

        # 1. Compute SHA-256 hash of the uploaded file contents
        sha256 = hashlib.sha256()
        for chunk in file_obj.chunks():
            sha256.update(chunk)
        file_hash = sha256.hexdigest()

        # 2. Check if this file has already been uploaded for this specific case
        same_case_duplicate = LegalDocument.objects.filter(
            case_file=case_file,
            content_hash=file_hash
        ).first()

        if same_case_duplicate:
            return Response(
                {"error": f"Document '{file_obj.name}' has already been uploaded for this case."},
                status=status.HTTP_409_CONFLICT
            )

        # 3. Check if file exists globally (uploaded by another case or lawyer)
        global_duplicate = LegalDocument.objects.filter(content_hash=file_hash).first()

        if global_duplicate:
            # Create a database record referencing the existing file path
            doc = LegalDocument.objects.create(
                case_file=case_file,
                file=global_duplicate.file,  # Points to the existing storage path
                name=file_obj.name,
                content_hash=file_hash,
                indexed=False
            )
        else:
            # Save the new file normally
            doc = LegalDocument.objects.create(
                case_file=case_file,
                file=file_obj,
                name=file_obj.name,
                content_hash=file_hash,
                indexed=False
            )

        # 4. Ingest file content into ChromaDB
        file_path = doc.file.path
        if os.path.exists(file_path):
            try:
                print(f"[Backend] Triggering ChromaDB ingestion for: {file_path}")
                process_legal_document(file_path, case_id, doc.id)
                doc.indexed = True
                doc.save()
            except Exception as e:
                doc.delete()  # Clean up the DB record if ingestion fails
                return Response(
                    {"error": f"Failed to ingest document: {str(e)}"},
                    status=status.HTTP_500_INTERNAL_SERVER_ERROR
                )

        serializer = LegalDocumentSerializer(doc)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

from io import BytesIO
import re
from django.http import HttpResponse
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_legal_docx(draft_text: str, title: str = "Legal Draft", doc_type: str = "bail", margins_inches: float = 1.0, font_family: str = "Times New Roman") -> BytesIO:
    """
    Generates a professional court/advocate standard .docx document with:
    - Configurable margins (default 1-inch)
    - Configurable font (default Times New Roman 12pt)
    - 1.5 line spacing
    - Centered bold court titles and headers
    - First-line indented numbered legal paragraphs (0.4 in)
    - Justified body paragraphs
    - Right-aligned signature and verification blocks
    """
    doc = docx.Document()
    
    # 1. Configurable Legal Margins
    section = doc.sections[0]
    section.top_margin = Inches(margins_inches)
    section.bottom_margin = Inches(margins_inches)
    section.left_margin = Inches(margins_inches)
    section.right_margin = Inches(margins_inches)
    
    # 2. Configure Normal Style
    normal_style = doc.styles['Normal']
    font = normal_style.font
    font.name = font_family
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    
    raw_lines = draft_text.split('\n')
    meta_label_pattern = re.compile(
        r'^(\d+\.\s*)?\*{0,2}(ADVOCATE\'?S?\s+LETTERHEAD.*|MODE OF TRANSMISSION.*|RECIPIENT PARTICULARS.*|SUBJECT LINE.*|SALUTATION & AUTHORIZATION STATEMENT.*|SALUTATION STATEMENT.*|STATEMENT OF FACTS & CAUSE OF ACTION.*|LEGAL GROUNDS & STATUTORY VIOLATIONS.*|FORMAL REQUISITION & DEMANDS.*|CONSEQUENCES OF NON-COMPLIANCE.*|RESERVATION OF RIGHTS & JURISDICTION.*|ADVOCATE SIGNATURE.*|COURT JURISDICTION|DOCUMENT TITLE|CAUSE TITLE)\*{0,2}\s*[:\-]?$',
        re.IGNORECASE
    )
    
    for i, raw_line in enumerate(raw_lines):
        line = raw_line.strip()
        if not line:
            continue
            
        # 1. Strip conversational preambles
        if re.match(r'^(here is|here\'s|below is|certainly|sure|please find|as requested|i will draft|i will generate|i shall draft|i have drafted)', line, re.IGNORECASE):
            continue

        # 2. Strip AI punctuation artifacts: wrapping parentheses '(' ')' or quotes
        line = re.sub(r'^\s*\(+[\'"]?', '', line)
        line = re.sub(r'[\'"]?\)+\s*$', '', line).strip()
        line = line.strip('\'"')

        # Clean markdown bold markers around labels
        clean_line = re.sub(r'^(\d+\.\s*)?\*{0,2}(COURT JURISDICTION|DOCUMENT TITLE|CAUSE TITLE|MEMO OF PARTIES)\*{0,2}\s*[:\-]\s*', '', line, flags=re.IGNORECASE)
        clean_line = clean_line.replace('**:', ':').replace(':**', ':')
        clean_line = re.sub(r'^[#*]+\s*', '', clean_line).rstrip('*# ')

        # Classify heading and alignment types
        is_court_header = any(keyword in clean_line.upper() for keyword in [
            'IN THE COURT OF', 'IN THE HIGH COURT', 'BEFORE THE HON\'BLE', 'IN THE SUPREME COURT'
        ])
        is_centered_title = any(keyword in clean_line.upper() for keyword in [
            'BAIL APPLICATION', 'LEGAL NOTICE', 'REPLY TO LEGAL NOTICE', 'AFFIDAVIT', 
            'NON-DISCLOSURE AGREEMENT', 'PETITION UNDER', 'APPLICATION UNDER', 
            'COMMERCIAL LEASE', 'CONSUMER COMPLAINT'
        ])
        is_memo_parties = bool(re.search(r'\.\.\.\s*(APPLICANT|PETITIONER|PLAINTIFF|COMPLAINANT)\s+VERSUS', clean_line, re.IGNORECASE))
        is_subject = bool(re.match(r'^(Sub|Subject)\s*[:\-]', clean_line, re.IGNORECASE))
        is_section_header = (
            clean_line.startswith('#') or 
            (line.startswith('**') and line.endswith('**') and len(line) < 100) or
            any(clean_line.upper().startswith(sec) for sec in [
                'ADVOCATE LETTERHEAD', 'MODE OF TRANSMISSION', 'NOTICEE PARTICULARS', 'SUBJECT LINE',
                'AUTHORIZATION STATEMENT', 'STATEMENT OF FACTS', 'FACTS OF THE CASE', 'GROUNDS', 
                'LEGAL GROUNDS', 'STATUTORY CITATIONS', 'DEMAND', 'DEMANDS', 'PRAYER', 'VERIFICATION', 
                'RECITALS', 'OPERATIVE', 'PRELIMINARY OBJECTIONS', 'PARA-WISE REPLY', 'ADVOCATE SIGNATURE',
                'COMMISSION JURISDICTION', 'DEPONENT IDENTIFICATION'
            ])
        )
        is_signature = any(clean_line.upper().startswith(sig) for sig in [
            'THROUGH:', 'COUNSEL FOR', 'DEPONENT', 'ADVOCATE', 'IN WITNESS WHEREOF', 
            'YOURS FAITHFULLY', '[ADVOCATE', 'ADV. '
        ])
        is_numbered_legal_para = bool(re.match(r'^(\d+\.)\s+', clean_line))

        is_letterhead = bool(re.search(r'^(Chambers of|Office of)', clean_line, re.IGNORECASE))

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(4)

        if is_court_header or is_centered_title:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(clean_line)
            run.bold = True
            run.font.size = Pt(13 if is_court_header else 14)
        elif is_memo_parties:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(clean_line)
            run.bold = True
            run.font.size = Pt(11)
        elif is_section_header:
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(clean_line)
            run.bold = True
            run.font.size = Pt(12)
        elif is_subject:
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(clean_line)
            run.bold = True
            run.font.size = Pt(12)
        elif is_letterhead:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean_line)
            run.font.size = Pt(11)
        elif is_signature:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_before = Pt(12)
            run = p.add_run(clean_line)
            run.bold = True
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if is_numbered_legal_para:
                p.paragraph_format.first_line_indent = Inches(0.4)
            # Handle inline markdown bolding (**bold text**)
            parts = re.split(r'(\*\*.*?\*\*)', clean_line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


class AIDraftGeneratorView(views.APIView):
    """
    POST /api/cases/draft/
    Triggers the LangGraph multi-agent loop to generate a Critic-approved legal draft.
    Supports document archetype selection and human-in-the-loop revision cycles.
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, *args, **kwargs):
        user_prompt = request.data.get('user_prompt')
        case_id = request.data.get('case_file')
        doc_type = request.data.get('doc_type', 'auto')
        user_feedback = request.data.get('user_feedback', '')
        previous_draft = request.data.get('previous_draft', '')
        revision_count = request.data.get('revision_count', 0)

        if not user_prompt:
            return Response(
                {"error": "Field 'user_prompt' is required."},
                status=status.HTTP_400_BAD_REQUEST
            )

        case_file = get_object_or_404(CaseFile, id=case_id, lawyer=request.user) if case_id else None

        print(f"[Backend] Starting LangGraph multi-agent system with prompt: {user_prompt} for Case ID: {case_id}, doc_type: {doc_type}")

        initial_state = {
            "user_prompt": user_prompt,
            "case_id": str(case_file.id) if case_file else "",
            "context_documents": "",
            "current_draft": previous_draft or "",
            "critic_feedback": "",
            "revision_count": int(revision_count) if revision_count else 0,
            "is_approved": False,
            "step_logs": [],
            "doc_type": doc_type or "auto",
            "user_feedback": user_feedback or "",
            "previous_draft": previous_draft or ""
        }

        try:
            # Execute the LangGraph loop (limit depth to 20 steps to allow enough drafting cycles)
            result = ai_app.invoke(initial_state, config={"recursion_limit": 20})
            
            return Response({
                "user_prompt": result.get("user_prompt"),
                "context_documents": result.get("context_documents"),
                "current_draft": result.get("current_draft"),
                "is_approved": result.get("is_approved"),
                "revision_count": result.get("revision_count"),
                "critic_feedback": result.get("critic_feedback"),
                "step_logs": result.get("step_logs", []),
                "doc_type": result.get("doc_type", doc_type)
            }, status=status.HTTP_200_OK)
            
        except Exception as e:
            return Response(
                {"error": f"LangGraph execution failed: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )


class LegalDraftDocxExportView(views.APIView):
    """
    POST /api/cases/draft/export-docx/
    Generates and streams a standardized professional .docx legal document.
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, *args, **kwargs):
        draft_text = request.data.get('draft_text', '')
        title = request.data.get('title', 'Legal_Draft')
        doc_type = request.data.get('doc_type', 'bail')
        case_id = request.data.get('case_id')
        margins = float(request.data.get('margin', 1.0))
        font_family = request.data.get('font_family', 'Times New Roman')

        if not draft_text:
            return Response(
                {"error": "Field 'draft_text' is required."},
                status=status.HTTP_400_BAD_REQUEST
            )

        sanitized_title = re.sub(r'[^a-zA-Z0-9_\-]', '_', title)
        filename = f"{sanitized_title}_{doc_type}.docx"

        buffer = create_legal_docx(
            draft_text, 
            title=title, 
            doc_type=doc_type, 
            margins_inches=margins, 
            font_family=font_family
        )

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


